"""
Busca de vagas via dorks (site:dominio "cargo" "local").

Modo interativo (pergunta tudo):
    jobsearch

Modo direto:
    jobsearch "backend engineer" -l remote -l latam            # geral
    jobsearch "backend engineer" -c br                         # só Brasil
    jobsearch "backend engineer" -c br -l remote --group ats   # Brasil + remoto
    jobsearch --list-countries
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sqlite3
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path

import requests

# ------------------------------------------------------------------ cores --
# ANSI puro, sem dependência. Desliga sozinho quando a saída não é um terminal
# (ex.: redirecionando para arquivo), com NO_COLOR definido ou com --no-color.
_CODES = {"bold": "1", "dim": "2", "underline": "4", "red": "31",
          "green": "32", "yellow": "33", "blue": "34", "magenta": "35", "cyan": "36"}
USE_COLOR = sys.stdout.isatty() and "NO_COLOR" not in os.environ
if os.name == "nt":
    os.system("")  # habilita ANSI no terminal do Windows


def paint(text: str, *styles: str) -> str:
    if not USE_COLOR or not styles:
        return text
    return "\033[" + ";".join(_CODES[x] for x in styles) + "m" + text + "\033[0m"


# ------------------------------------------------------------------ logging --

log = logging.getLogger("jobsearch")
QUIET = False


def setup_logging(log_file: str | None = None, quiet: bool = False,
                  verbose: bool = False) -> None:
    """Configura logging para stdout/stderr e, opcionalmente, arquivo."""
    global QUIET
    QUIET = quiet
    level = logging.DEBUG if verbose else logging.INFO
    log.setLevel(level)
    log.handlers.clear()

    fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s",
                            datefmt="%Y-%m-%d %H:%M:%S")

    stream = logging.StreamHandler(sys.stderr)
    stream.setFormatter(fmt)
    stream.setLevel(logging.WARNING if quiet else level)
    log.addHandler(stream)

    if log_file:
        fileh = logging.FileHandler(log_file, encoding="utf-8")
        fileh.setFormatter(fmt)
        fileh.setLevel(level)
        log.addHandler(fileh)

    log.propagate = False


# ------------------------------------------------------------------- datas --

_RELATIVE_PATTERNS = [
    (re.compile(r"\b(\d+)\s*(?:minute|min)s?\s+ago\b", re.I), "minutes"),
    (re.compile(r"\b(\d+)\s*(?:hour|hr)s?\s+ago\b", re.I), "hours"),
    (re.compile(r"\b(\d+)\s*days?\s+ago\b", re.I), "days"),
    (re.compile(r"\b(\d+)\s*(?:week|wk)s?\s+ago\b", re.I), "weeks"),
    (re.compile(r"\b(\d+)\s*months?\s+ago\b", re.I), "months"),
    (re.compile(r"\b(\d+)\s*(?:minuto|min)s?\s+atr[áa]s\b", re.I), "minutes"),
    (re.compile(r"\b(\d+)\s*horas?\s+atr[áa]s\b", re.I), "hours"),
    (re.compile(r"\b(\d+)\s*dias?\s+atr[áa]s\b", re.I), "days"),
    (re.compile(r"\b(\d+)\s*semanas?\s+atr[áa]s\b", re.I), "weeks"),
    (re.compile(r"\b(\d+)\s*meses?\s+atr[áa]s\b", re.I), "months"),
    # "há 3 dias", "ha 2 horas"
    (re.compile(r"\bh[áa]\s+(\d+)\s*(?:minuto|min)s?\b", re.I), "minutes"),
    (re.compile(r"\bh[áa]\s+(\d+)\s*horas?\b", re.I), "hours"),
    (re.compile(r"\bh[áa]\s+(\d+)\s*dias?\b", re.I), "days"),
    (re.compile(r"\bh[áa]\s+(\d+)\s*semanas?\b", re.I), "weeks"),
    (re.compile(r"\bh[áa]\s+(\d+)\s*meses?\b", re.I), "months"),
]


def parse_posted_date(value, now: datetime | None = None) -> datetime | None:
    """
    Extrai uma data de publicação a partir de formatos comuns.

    Aceita datetime, número (dias atrás), ISO ("2026-09-28"), texto relativo
    ("2 days ago", "há 3 dias", "today", "yesterday").
    """
    if value is None:
        return None
    now = now or datetime.now()

    if isinstance(value, datetime):
        return value
    if isinstance(value, (int, float)):
        # convenção: número = dias atrás
        return now - timedelta(days=float(value))

    text = str(value).strip()
    if not text:
        return None

    low = text.lower()
    if low in ("today", "hoje"):
        return now
    if low in ("yesterday", "ontem"):
        return now - timedelta(days=1)

    for pattern, unit in _RELATIVE_PATTERNS:
        m = pattern.search(text)
        if m:
            amount = int(m.group(1))
            if unit == "minutes":
                return now - timedelta(minutes=amount)
            if unit == "hours":
                return now - timedelta(hours=amount)
            if unit == "days":
                return now - timedelta(days=amount)
            if unit == "weeks":
                return now - timedelta(weeks=amount)
            if unit == "months":
                return now - timedelta(days=30 * amount)

    # ISO / dateutil-like "YYYY-MM-DD[ HH:MM:SS]"
    candidate = text.replace("T", " ").split(".")[0]
    candidate = re.sub(r"(Z|[+-]\d{2}:?\d{2})$", "", candidate).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
                "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(candidate, fmt)
        except ValueError:
            continue
    return None


# ----------------------------------------------------------- enriquecimento --

# Padrões que expõem o nome da empresa em URLs de ATS/portais.
_COMPANY_URL_PATTERNS = [
    re.compile(r"jobs\.lever\.co/([^/?#]+)", re.I),
    re.compile(r"(?:boards|job-boards)\.greenhouse\.io/([^/?#]+)", re.I),
    re.compile(r"apply\.workable\.com/([^/?#]+)", re.I),
    re.compile(r"jobs\.ashbyhq\.com/([^/?#]+)", re.I),
    re.compile(r"smartrecruiters\.com/([^/?#]+)", re.I),
    re.compile(r"indeed\.[a-z.]+/cmp/([^/?#]+)", re.I),
    re.compile(r"glassdoor\.[a-z.]+/(?:Overview|Jobs)/[^/]*?EI_IE\d+\.\d+,\d+_([^/?#]+)", re.I),
]

# Títulos costumam vir como "Vaga na Empresa", "Vaga - Empresa", "Vaga | Empresa"
_TITLE_COMPANY_PATTERNS = [
    re.compile(r"\s+(?:at|@)\s+([A-Z][\w&.\- ]{1,40})$"),
    re.compile(r"\s+[-–|]\s+([A-Z][\w&.\- ]{1,40})$"),
]

_SUBDOMAIN_ATS = [
    re.compile(r"^(?P<company>[^.]+)\.gupy\.io$", re.I),
    re.compile(r"^(?P<company>[^.]+)\.(?:wd\d+\.)?myworkdayjobs\.com$", re.I),
]


def _pretty_company(token: str) -> str:
    """Converte slug/domínio em nome legível: 'acme-corp' -> 'Acme Corp'."""
    token = token.strip()
    token = re.sub(r"\.(com|io|co|jobs|net|org).*$", "", token, flags=re.I)
    token = re.sub(r"[-_+]+", " ", token)
    token = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", token)
    token = re.sub(r"\s+", " ", token).strip()
    if not token:
        return ""
    if token.islower() or token.isupper():
        return token.title()
    return token


def extract_company(link: str = "", title: str = "") -> str:
    """
    Extrai o nome da empresa a partir do link (ATS/portais) ou do título.

    Retorna string vazia quando não for possível identificar.
    """
    from urllib.parse import urlsplit

    if link:
        try:
            host = urlsplit(link).netloc.lower().split(":")[0]
        except ValueError:
            host = ""
        for pattern in _SUBDOMAIN_ATS:
            m = pattern.match(host)
            if m:
                return _pretty_company(m.group("company"))
        for pattern in _COMPANY_URL_PATTERNS:
            m = pattern.search(link)
            if m:
                return _pretty_company(m.group(1))

    if title:
        clean = title.strip()
        for pattern in _TITLE_COMPANY_PATTERNS:
            m = pattern.search(clean)
            if m:
                return _pretty_company(m.group(1))
    return ""


# Portais genéricos (funcionam para busca geral)
BOARDS = ["indeed.com", "linkedin.com/jobs", "glassdoor.com"]

# ATS usados pelas empresas (ouro escondido)
ATS = [
    "boards.greenhouse.io",
    "jobs.lever.co",
    "myworkdayjobs.com",
    "smartrecruiters.com",
    "jobs.ashbyhq.com",
    "apply.workable.com",
]

GROUPS = {"boards": BOARDS, "ats": ATS, "all": BOARDS + ATS}

# Config por país. Edite à vontade: "names" são os termos usados na query
# (com OR entre eles), "region" vai para o backend, "indeed"/"glassdoor" são
# os domínios locais e "extra" são portais que só existem naquele país.
COUNTRIES = {
    "br": {"names": ["Brazil", "Brasil"], "region": "br-pt", "gl": "br",
           "indeed": "br.indeed.com", "glassdoor": "glassdoor.com.br",
           "extra": ["gupy.io", "infojobs.com.br", "vagas.com.br"]},
    "pt": {"names": ["Portugal"], "region": "pt-pt", "gl": "pt",
           "indeed": "pt.indeed.com", "extra": []},
    "us": {"names": ["United States", "USA"], "region": "us-en", "gl": "us",
           "indeed": "indeed.com", "extra": []},
    "uk": {"names": ["United Kingdom", "UK"], "region": "uk-en", "gl": "uk",
           "indeed": "uk.indeed.com", "glassdoor": "glassdoor.co.uk", "extra": []},
    "ca": {"names": ["Canada"], "region": "ca-en", "gl": "ca",
           "indeed": "ca.indeed.com", "glassdoor": "glassdoor.ca", "extra": []},
    "de": {"names": ["Germany", "Deutschland"], "region": "de-de", "gl": "de",
           "indeed": "de.indeed.com", "glassdoor": "glassdoor.de", "extra": []},
    "es": {"names": ["Spain", "España"], "region": "es-es", "gl": "es",
           "indeed": "es.indeed.com", "glassdoor": "glassdoor.es", "extra": []},
    "mx": {"names": ["Mexico", "México"], "region": "mx-es", "gl": "mx",
           "indeed": "mx.indeed.com", "glassdoor": "glassdoor.com.mx", "extra": []},
    "ar": {"names": ["Argentina"], "region": "ar-es", "gl": "ar",
           "indeed": "ar.indeed.com", "glassdoor": "glassdoor.com.ar", "extra": []},
    "co": {"names": ["Colombia"], "region": "co-es", "gl": "co",
           "indeed": "co.indeed.com", "extra": []},
    "cl": {"names": ["Chile"], "region": "cl-es", "gl": "cl",
           "indeed": "cl.indeed.com", "extra": []},
}


# ----------------------------------------------------------------- queries --

def country_term(country: dict) -> str:
    names = [f'"{n}"' for n in country["names"]]
    return names[0] if len(names) == 1 else "(" + " OR ".join(names) + ")"


def domains_for(group: str, country: dict | None) -> list[str]:
    domains = list(GROUPS[group])
    if not country:
        return domains
    out = []
    for d in domains:
        if d == "indeed.com":
            d = country["indeed"]
        elif d == "glassdoor.com":
            d = country.get("glassdoor", "glassdoor.com")
        out.append(d)
    if group in ("boards", "all"):
        out += country["extra"]
    return out


def plan_queries(role: str, locals_: list[str], group: str,
                 country: dict | None) -> list[dict]:
    """Retorna [{domain, location, query}]."""
    plan = []
    for d in domains_for(group, country):
        if country:
            for r in (locals_ or [None]):
                terms = [country_term(country)] + ([f'"{r}"'] if r else [])
                label = country["names"][0] + (f" / {r}" if r else "")
                plan.append({"domain": d, "location": label,
                             "query": f'site:{d} "{role}" ' + " ".join(terms)})
        else:
            for loc in (locals_ or ["remote"]):
                plan.append({"domain": d, "location": loc,
                             "query": f'site:{d} "{role}" "{loc}"'})
    return plan


# ---------------------------------------------------------------- backends --

class SearchError(Exception):
    """Erro durante a busca."""
    pass


def search_ddg(query: str, max_results: int, country: dict | None,
               retries: int = 3, backoff: float = 2.0) -> list[dict]:
    """Busca no DuckDuckGo com retry automático."""
    try:
        from ddgs import DDGS
    except ImportError as e:
        raise SearchError("ddgs não instalado. Execute: pip install ddgs") from e

    kwargs = {"max_results": max_results}
    if country:
        kwargs["region"] = country["region"]

    last_error = None
    for attempt in range(retries):
        try:
            results = DDGS().text(query, **kwargs)
            return [{"title": r.get("title", ""), "link": r.get("href", ""),
                     "snippet": r.get("body", "")}
                    for r in results]
        except Exception as e:
            last_error = e
            if attempt < retries - 1:
                wait = backoff * (2 ** attempt)
                print(paint(f"   [retry {attempt + 1}/{retries}] erro: {e}. "
                            f"Tentando novamente em {wait:.1f}s...", "yellow"))
                time.sleep(wait)

    raise SearchError(f"Falha após {retries} tentativas: {last_error}")


def search_google(query: str, max_results: int, country: dict | None,
                  retries: int = 3, backoff: float = 2.0) -> list[dict]:
    """Busca no Google Custom Search com retry automático."""
    key = os.environ.get("GOOGLE_API_KEY")
    cx = os.environ.get("GOOGLE_CX")
    if not key or not cx:
        raise SearchError("GOOGLE_API_KEY e GOOGLE_CX devem ser definidos para usar o backend google")

    params = {"key": key, "cx": cx, "q": query, "num": min(max_results, 10)}
    if country:
        params["gl"] = country["gl"]

    last_error = None
    for attempt in range(retries):
        try:
            resp = requests.get("https://www.googleapis.com/customsearch/v1",
                                params=params, timeout=15)
            resp.raise_for_status()
            return [{"title": i.get("title", ""), "link": i.get("link", ""),
                     "snippet": i.get("snippet", ""),
                     "posted_at": _google_posted_date(i)}
                    for i in resp.json().get("items", [])]
        except requests.exceptions.RequestException as e:
            last_error = e
            if attempt < retries - 1:
                wait = backoff * (2 ** attempt)
                print(paint(f"   [retry {attempt + 1}/{retries}] erro: {e}. "
                            f"Tentando novamente em {wait:.1f}s...", "yellow"))
                time.sleep(wait)

    raise SearchError(f"Falha após {retries} tentativas: {last_error}")


def _google_posted_date(item: dict) -> str | None:
    """Tenta extrair a data de publicação do pagemap do Google CSE."""
    metatags = (item.get("pagemap") or {}).get("metatags") or []
    keys = ("article:published_time", "datepublished", "date", "og:updated_time",
            "article:modified_time", "pubdate")
    for mt in metatags:
        for k in keys:
            if k in mt and mt[k]:
                return mt[k]
    return None


BACKENDS = {"ddg": search_ddg, "google": search_google}


# ------------------------------------------------------------------ filtros --

def filter_jobs(jobs: list[dict], min_date: str | None = None,
                max_date: str | None = None,
                keywords: list[str] | None = None,
                exclude_keywords: list[str] | None = None,
                max_days: int = 0,
                keep_unknown_dates: bool = True) -> list[dict]:
    """
    Filtra vagas por data e palavras-chave.

    Args:
        jobs: Lista de vagas
        min_date: Data mínima (formato: YYYY-MM-DD)
        max_date: Data máxima (formato: YYYY-MM-DD)
        keywords: Palavras-chave que devem estar no título
        exclude_keywords: Palavras-chave que NÃO devem estar no título
        max_days: Idade máxima em dias (0 desativa)
        keep_unknown_dates: manter vagas sem data identificável

    Returns:
        Lista de vagas filtradas
    """
    filtered = jobs.copy()

    if keywords:
        filtered = [j for j in filtered
                    if any(kw.lower() in j["title"].lower() for kw in keywords)]

    if exclude_keywords:
        filtered = [j for j in filtered
                    if not any(kw.lower() in j["title"].lower() for kw in exclude_keywords)]

    if max_days and max_days > 0:
        cutoff = datetime.now() - timedelta(days=max_days)
        filtered_new = []
        for j in filtered:
            posted = _job_posted_datetime(j)
            if posted is None:
                if keep_unknown_dates:
                    filtered_new.append(j)
                continue
            if posted >= cutoff:
                filtered_new.append(j)
        filtered = filtered_new

    if min_date:
        try:
            lo = datetime.strptime(min_date, "%Y-%m-%d")
            filtered = [j for j in filtered
                        if _in_date_range(j, lo, None, keep_unknown_dates)]
        except ValueError:
            log.warning("min_date inválida (use YYYY-MM-DD): %s", min_date)

    if max_date:
        try:
            hi = datetime.strptime(max_date, "%Y-%m-%d") + timedelta(days=1)
            filtered = [j for j in filtered
                        if _in_date_range(j, None, hi, keep_unknown_dates)]
        except ValueError:
            log.warning("max_date inválida (use YYYY-MM-DD): %s", max_date)

    return filtered


def _in_date_range(job: dict, lo: datetime | None, hi: datetime | None,
                   keep_unknown: bool) -> bool:
    """Verifica se a vaga está no intervalo [lo, hi). Sem data, segue keep_unknown."""
    posted = _job_posted_datetime(job)
    if posted is None:
        return keep_unknown
    if lo is not None and posted < lo:
        return False
    if hi is not None and posted >= hi:
        return False
    return True


def _job_posted_datetime(job: dict) -> datetime | None:
    """Retorna a data de publicação da vaga, se identificável."""
    for key in ("posted_at", "date", "published", "age_days"):
        value = job.get(key)
        if value is None:
            continue
        dt = parse_posted_date(value)
        if dt is not None:
            return dt
    # tenta extrair de snippet/body ("2 days ago")
    for key in ("snippet", "body"):
        text = job.get(key)
        if text:
            dt = parse_posted_date(text)
            if dt is not None:
                return dt
    return None


# ------------------------------------------------------------------- sqlite --

_TRACKING_PARAMS = re.compile(
    r"^(utm_|gclid|fbclid|mc_|ref$|referrer$|source$|trk$|tracking)", re.I)


def normalize_url(url: str) -> str:
    """Normaliza URL para deduplicação: remove tracking, fragmento e barra final."""
    if not url:
        return url
    from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
    try:
        parts = urlsplit(url.strip())
    except ValueError:
        return url.strip()
    query = [(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True)
             if not _TRACKING_PARAMS.match(k)]
    path = parts.path.rstrip("/") or "/"
    netloc = parts.netloc.lower()
    return urlunsplit((parts.scheme.lower(), netloc, path,
                       urlencode(query), ""))


def _get_db_path(db: str | None) -> Path:
    if db:
        return Path(db).expanduser().resolve()
    return Path.cwd() / "jobsearch.db"


def _connect(db: str | None) -> sqlite3.Connection:
    db_path = _get_db_path(db)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def _ensure_schema(conn: sqlite3.Connection) -> None:
    cur = conn.cursor()
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            link TEXT NOT NULL UNIQUE,
            link_key TEXT,
            source TEXT,
            domain TEXT,
            location TEXT,
            query TEXT,
            posted_at TEXT,
            salary TEXT,
            company TEXT,
            description TEXT,
            first_seen_at TEXT,
            last_seen_at TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    # Migração para bancos criados por versões anteriores.
    cols = {row[1] for row in cur.execute("PRAGMA table_info(jobs)")}
    for col in ("link_key", "first_seen_at", "last_seen_at"):
        if col not in cols:
            cur.execute(f"ALTER TABLE jobs ADD COLUMN {col} TEXT")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_jobs_link_key ON jobs(link_key)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_jobs_created_at ON jobs(created_at)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_jobs_source ON jobs(source)")
    conn.commit()


def save_to_db(jobs: list[dict], db: str | None = None) -> int:
    """
    Salva vagas em banco SQLite.

    Deduplica pela URL normalizada (sem tracking). Retorna o número de vagas
    novas inseridas. Vagas já existentes têm `last_seen_at` atualizado.
    """
    if not jobs:
        return 0

    conn = _connect(db)
    try:
        _ensure_schema(conn)
        cur = conn.cursor()
        now = datetime.now().isoformat(timespec="seconds")
        inserted = 0
        for j in jobs:
            link = j.get("link", "")
            if not link:
                continue
            try:
                cur.execute(
                    """
                    INSERT OR IGNORE INTO jobs (
                        title, link, link_key, source, domain, location, query,
                        posted_at, salary, company, description,
                        first_seen_at, last_seen_at, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        j.get("title", ""),
                        link,
                        normalize_url(link),
                        j.get("source") or j.get("domain"),
                        j.get("domain") or j.get("source"),
                        j.get("location"),
                        j.get("query"),
                        j.get("posted_at") or j.get("date") or j.get("published"),
                        j.get("salary"),
                        j.get("company"),
                        j.get("snippet") or j.get("description"),
                        now,
                        now,
                        now,
                        now,
                    ),
                )
                if cur.rowcount > 0:
                    inserted += 1
                else:
                    cur.execute(
                        "UPDATE jobs SET last_seen_at = ?, updated_at = ? WHERE link_key = ?",
                        (now, now, normalize_url(link)),
                    )
            except sqlite3.Error as e:
                log.warning("Falha ao salvar vaga no banco: %s", e)
                continue
        conn.commit()
    finally:
        conn.close()

    return inserted


def db_stats(db: str | None = None) -> dict:
    """Retorna estatísticas do banco."""
    conn = _connect(db)
    try:
        _ensure_schema(conn)
        cur = conn.cursor()
        total = cur.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
        first = cur.execute("SELECT MIN(created_at) FROM jobs").fetchone()[0]
        last = cur.execute("SELECT MAX(created_at) FROM jobs").fetchone()[0]
        by_source = cur.execute(
            "SELECT COALESCE(source, domain) AS s, COUNT(*) c FROM jobs "
            "GROUP BY s ORDER BY c DESC"
        ).fetchall()
        by_day = cur.execute(
            "SELECT substr(created_at, 1, 10) AS d, COUNT(*) c FROM jobs "
            "GROUP BY d ORDER BY d DESC LIMIT 14"
        ).fetchall()
    finally:
        conn.close()
    return {
        "total": total,
        "first_seen": first,
        "last_seen": last,
        "by_source": [(r["s"], r["c"]) for r in by_source],
        "by_day": [(r["d"], r["c"]) for r in by_day],
    }


def db_export(db: str | None = None, fmt: str = "json", limit: int | None = None,
              since_days: int | None = None) -> str:
    """Exporta vagas do banco como JSON ou CSV (retorna o texto)."""
    import csv
    import io

    conn = _connect(db)
    try:
        _ensure_schema(conn)
        sql = "SELECT * FROM jobs"
        params: list = []
        if since_days:
            cutoff = (datetime.now() - timedelta(days=since_days)).isoformat(timespec="seconds")
            sql += " WHERE created_at >= ?"
            params.append(cutoff)
        sql += " ORDER BY created_at DESC"
        if limit:
            sql += " LIMIT ?"
            params.append(limit)
        rows = [dict(r) for r in conn.execute(sql, params)]
    finally:
        conn.close()

    if fmt == "csv":
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        return buf.getvalue()
    return json.dumps(rows, ensure_ascii=False, indent=2)


def db_purge(db: str | None = None, older_than_days: int | None = None) -> int:
    """Remove vagas antigas. Sem `older_than_days`, limpa todo o banco."""
    conn = _connect(db)
    try:
        _ensure_schema(conn)
        cur = conn.cursor()
        if older_than_days and older_than_days > 0:
            cutoff = (datetime.now() - timedelta(days=older_than_days)).isoformat(timespec="seconds")
            cur.execute("DELETE FROM jobs WHERE created_at < ?", (cutoff,))
        else:
            cur.execute("DELETE FROM jobs")
        deleted = cur.rowcount
        conn.commit()
    finally:
        conn.close()
    return deleted


# ------------------------------------------------------------- interactive --

def ask(prompt: str, default: str = "") -> str:
    suffix = f" [{default}]" if default else ""
    answer = input(f"{prompt}{suffix}: ").strip()
    return answer or default


def default_basename(role: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", role.lower()).strip("-") or "vagas"
    return f"vagas_{slug}_{datetime.now():%Y%m%d_%H%M}"


def interactive_prompts(args: argparse.Namespace) -> None:
    print(paint("=== Busca de vagas ===", "bold", "cyan"), "\n")
    while not args.role:
        args.role = ask("Qual vaga você está buscando? (ex: backend engineer)")

    print(paint("\nAlcance da busca:", "bold"))
    print("  1) Geral (mundo todo)")
    print("  2) Um país específico")
    if ask("Opção", "1") == "2":
        print("  Países: " + ", ".join(f"{k} ({v['names'][0]})" for k, v in COUNTRIES.items()))
        code = ask("Código do país", "br").lower()
        if code in COUNTRIES:
            args.country = code
        else:
            print(paint(f"  [aviso] país '{code}' desconhecido, usando busca geral.", "yellow"))
        locs = ask("Refinar por cidade/remote (opcional, vírgula)", "")
    else:
        args.country = None
        locs = ask("Local(is), separados por vírgula", "remote, latam")
    args.local = [x.strip() for x in locs.split(",") if x.strip()]

    print(paint("\nOnde buscar?", "bold"))
    print("  1) Portais (Indeed, LinkedIn, Glassdoor + locais do país)")
    print("  2) ATS das empresas (Greenhouse, Lever, Workday, ...)")
    print("  3) Todos")
    args.group = {"1": "boards", "2": "ats", "3": "all"}.get(ask("Opção", "3"), "all")

    max_raw = ask("Resultados por query", str(args.max))
    args.max = int(max_raw) if max_raw.isdigit() else args.max

    # Filtros adicionais
    print(paint("\nFiltros adicionais (opcional):", "bold"))
    include = ask("Palavras-chave incluir (vírgula)", "")
    args.filter_include = [x.strip() for x in include.split(",") if x.strip()]
    exclude = ask("Palavras-chave excluir (vírgula)", "")
    args.filter_exclude = [x.strip() for x in exclude.split(",") if x.strip()]

    max_days_raw = ask("Idade máxima da vaga em dias (0 desativa)", str(args.max_days))
    if max_days_raw.isdigit():
        args.max_days = int(max_days_raw)

    args.output = ask("Nome base dos arquivos (gera .json e .xlsx)",
                      default_basename(args.role))

    print(paint("\nSalvar em banco SQLite?", "bold"))
    db_path = ask("Caminho do banco (vazio = não)", "")
    args.db = db_path or None

    args.csv = ask("Também salvar CSV? (s/N)", "n").lower() in ("s", "sim", "y", "yes")
    print()


# ------------------------------------------------------------------ output --

def save_json(path: str, meta: dict, jobs: list[dict]) -> None:
    payload = {**meta, "total": len(jobs), "jobs": jobs}
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                          encoding="utf-8")


def save_xlsx(path: str, meta: dict, jobs: list[dict]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(paint("[aviso] openpyxl não instalado (pip install openpyxl); .xlsx não foi gerado.", "yellow"))
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Vagas"

    headers = ["#", "Título", "Empresa", "Link", "Domínio", "Local", "Query", "Descrição"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center")

    for i, j in enumerate(jobs, start=1):
        ws.append([i, j.get("title", ""), j.get("company", ""), j.get("link", ""),
                   j.get("domain", ""), j.get("location", ""), j.get("query", ""),
                   j.get("description", "")])
        link_cell = ws.cell(row=i + 1, column=4)
        link_cell.hyperlink = j.get("link", "")
        link_cell.font = Font(color="0563C1", underline="single")

    # Larguras dinâmicas com base no conteúdo (com limites razoáveis).
    caps = (5, 55, 28, 60, 22, 20, 45, 70)
    for idx, cap in enumerate(caps, start=1):
        longest = len(str(headers[idx - 1]))
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx,
                                values_only=True):
            if row[0]:
                longest = max(longest, len(str(row[0])))
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, cap)
    ws.freeze_panes = "A2"
    if jobs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(jobs) + 1}"

    info = wb.create_sheet("Busca")
    for k, v in meta.items():
        info.append([k, ", ".join(v) if isinstance(v, list) else v])
    info.append(["total", len(jobs)])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 40

    wb.save(path)
    return True


_CSV_FIELDS = ["title", "company", "link", "domain", "location", "query",
               "posted_at", "description"]


def save_csv(path: str, jobs: list[dict]) -> bool:
    """Salva as vagas em CSV (UTF-8). Retorna True em caso de sucesso."""
    import csv

    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=_CSV_FIELDS, extrasaction="ignore")
            writer.writeheader()
            for j in jobs:
                writer.writerow({k: j.get(k, "") or "" for k in _CSV_FIELDS})
    except OSError as e:
        print(paint(f"[aviso] não foi possível gerar o CSV: {e}", "yellow"))
        return False
    return True


# ------------------------------------------------------------- db command --

def db_command(argv: list[str]) -> int:
    """Subcomandos para consultar/gerenciar o banco: stats, export, purge."""
    parser = argparse.ArgumentParser(prog="jobsearch db",
                                     description="Utilitários do banco SQLite")
    parser.add_argument("--db", help="caminho do banco (padrão: ./jobsearch.db)")
    sub = parser.add_subparsers(dest="action", required=True)

    p_stats = sub.add_parser("stats", help="mostra estatísticas do banco")
    p_stats.add_argument("--json", action="store_true", help="saída em JSON")

    p_export = sub.add_parser("export", help="exporta as vagas")
    p_export.add_argument("-o", "--output", help="arquivo de saída (padrão: stdout)")
    p_export.add_argument("--format", choices=["json", "csv"], default="json")
    p_export.add_argument("--limit", type=int, help="máximo de vagas")
    p_export.add_argument("--since-days", type=int, help="somente dos últimos N dias")

    p_purge = sub.add_parser("purge", help="remove vagas antigas")
    p_purge.add_argument("--older-than", type=int, metavar="DIAS",
                         help="remove vagas vistas há mais de N dias")
    p_purge.add_argument("-y", "--yes", action="store_true",
                         help="não pedir confirmação")

    args = parser.parse_args(argv)
    db = args.db

    if args.action == "stats":
        stats = db_stats(db)
        if args.json:
            print(json.dumps(stats, ensure_ascii=False, indent=2))
            return 0
        print(paint("Banco:", "bold"), _get_db_path(db))
        print(paint("Total de vagas:", "bold"), stats["total"])
        print("Primeira:", stats["first_seen"] or "-")
        print("Última:  ", stats["last_seen"] or "-")
        if stats["by_source"]:
            print("\n" + paint("Por fonte:", "bold"))
            for src, count in stats["by_source"]:
                print(f"  {count:>5}  {src}")
        if stats["by_day"]:
            print("\n" + paint("Por dia (últimos):", "bold"))
            for day, count in stats["by_day"]:
                print(f"  {count:>5}  {day}")
        return 0

    if args.action == "export":
        content = db_export(db, fmt=args.format, limit=args.limit,
                            since_days=args.since_days)
        if args.output:
            Path(args.output).write_text(content, encoding="utf-8")
            print(f"Exportado para {args.output}")
        else:
            print(content)
        return 0

    if args.action == "purge":
        if not args.yes:
            target = (f"vagas mais antigas que {args.older_than} dias"
                      if args.older_than else "TODAS as vagas")
            resp = input(f"Confirmar remoção de {target}? [s/N]: ").strip().lower()
            if resp not in ("s", "y", "sim", "yes"):
                print("Cancelado.")
                return 1
        deleted = db_purge(db, older_than_days=args.older_than)
        print(f"{deleted} vaga(s) removida(s).")
        return 0

    return 1  # pragma: no cover


# -------------------------------------------------------------------- main --

def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Busca de vagas com dorks de busca")
    p.add_argument("role", nargs="?", help='cargo, ex: "backend engineer" (omita para modo interativo)')
    p.add_argument("-c", "--country", choices=COUNTRIES,
                   help="limita a busca a um país (omita para busca geral)")
    p.add_argument("-l", "--local", action="append", default=[],
                   help="local/termo extra (repita para vários). Com -c refina o país (cidade, remote)")
    p.add_argument("-g", "--group", choices=GROUPS, default="all")
    p.add_argument("-b", "--backend", choices=BACKENDS, default="ddg")
    p.add_argument("-m", "--max", type=int, default=8, help="resultados por query")
    p.add_argument("-o", "--output", help="nome base dos arquivos de saída, gera .json e .xlsx")
    p.add_argument("-i", "--interactive", action="store_true", help="força o modo interativo")
    p.add_argument("--delay", type=float, default=2.0, help="segundos entre queries")
    p.add_argument("--retries", type=int, default=3, help="número de tentativas em caso de erro")
    p.add_argument("--show-queries", action="store_true", help="só imprime as queries")
    p.add_argument("--no-color", action="store_true", help="desativa as cores")
    p.add_argument("--list-countries", action="store_true", help="lista os países disponíveis")
    p.add_argument("--filter-include", nargs="+", default=[],
                   help="palavras-chave que devem estar no título")
    p.add_argument("--filter-exclude", nargs="+", default=[],
                   help="palavras-chave que NÃO devem estar no título")
    p.add_argument("--max-days", type=int, default=3,
                   help="idade máxima da vaga em dias (0 desativa; padrão: 3)")
    p.add_argument("--strict-dates", action="store_true",
                   help="descarta vagas sem data identificável")
    p.add_argument("--min-date", help="data mínima de publicação (YYYY-MM-DD)")
    p.add_argument("--max-date", help="data máxima de publicação (YYYY-MM-DD)")
    p.add_argument("--db", help="salvar resultados em banco SQLite (ex.: jobsearch.db)")
    p.add_argument("--csv", action="store_true", help="também salvar arquivo CSV")
    p.add_argument("--no-json", action="store_true", help="não salvar arquivo JSON")
    p.add_argument("--no-xlsx", action="store_true", help="não salvar arquivo XLSX")
    p.add_argument("-q", "--quiet", action="store_true", help="suprime saída de progresso")
    p.add_argument("-v", "--verbose", action="store_true", help="log detalhado")
    p.add_argument("--log-file", help="arquivo de log (append)")
    return p


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)

    # Subcomandos de banco: jobsearch db <stats|export|purge>
    if argv and argv[0] == "db":
        return db_command(argv[1:])

    p = _build_parser()
    args = p.parse_args(argv)

    global USE_COLOR
    if args.no_color:
        USE_COLOR = False

    setup_logging(log_file=args.log_file, quiet=args.quiet, verbose=args.verbose)

    if args.list_countries:
        for k, v in COUNTRIES.items():
            print(f"{paint(k, 'bold', 'cyan')}  {v['names'][0]}")
        return 0

    if args.interactive or not args.role:
        interactive_prompts(args)

    country = COUNTRIES.get(args.country) if args.country else None
    plan = plan_queries(args.role, args.local, args.group, country)

    if args.show_queries:
        for item in plan:
            print(item["query"])
        return 0

    scope = country["names"][0] if country else "geral"
    log.info("Escopo: %s | %d queries", scope, len(plan))

    search = BACKENDS[args.backend]
    seen: set[str] = set()
    jobs: list[dict] = []
    errors = 0

    try:
        for item in plan:
            if not QUIET:
                print("\n" + paint(f"== {item['domain']}", "bold", "cyan"),
                      paint(f"| {item['location']} ==", "yellow"))
                print("   " + paint(item["query"], "dim"))
            try:
                results = search(item["query"], args.max, country,
                                 retries=args.retries)
            except SearchError as e:
                log.error("%s", e)
                errors += 1
                time.sleep(args.delay)
                continue

            new = [r for r in results if r["link"] and r["link"] not in seen]
            if not new and not QUIET:
                print(paint("   (nenhum resultado novo)", "dim"))
            for r in new:
                seen.add(r["link"])
                snippet = r.get("snippet", "") or ""
                jobs.append({"title": r["title"], "link": r["link"],
                             "domain": item["domain"], "location": item["location"],
                             "query": item["query"],
                             "source": item["domain"],
                             "company": extract_company(r["link"], r["title"]),
                             "description": snippet,
                             "snippet": snippet,
                             "posted_at": r.get("posted_at") or r.get("date")
                             or r.get("published")})
                if not QUIET:
                    print(f"   {paint('-', 'green')} {paint(r['title'], 'bold', 'green')}")
                    print("     " + paint(r["link"], "blue", "underline"))
            time.sleep(args.delay)
    except KeyboardInterrupt:
        log.warning("interrompido; salvando o que foi encontrado até aqui...")

    # Aplicar filtros (sempre que qualquer critério estiver ativo)
    has_filter = (args.filter_include or args.filter_exclude
                  or args.max_days > 0 or args.min_date or args.max_date)
    if has_filter:
        before = len(jobs)
        jobs = filter_jobs(
            jobs,
            keywords=args.filter_include,
            exclude_keywords=args.filter_exclude,
            max_days=args.max_days,
            min_date=args.min_date,
            max_date=args.max_date,
            keep_unknown_dates=not args.strict_dates,
        )
        log.info("Filtros aplicados: %d -> %d vagas", before, len(jobs))

    base = args.output or default_basename(args.role)
    base = re.sub(r"\.(json|xlsx?)$", "", base, flags=re.IGNORECASE)
    meta = {
        "role": args.role,
        "scope": scope,
        "country": args.country or "",
        "locations": args.local,
        "group": args.group,
        "backend": args.backend,
        "filter_include": args.filter_include,
        "filter_exclude": args.filter_exclude,
        "max_days": args.max_days,
        "searched_at": datetime.now().isoformat(timespec="seconds"),
    }
    saved_list = []
    if not args.no_json:
        save_json(f"{base}.json", meta, jobs)
        saved_list.append(f"{base}.json")
    if not args.no_xlsx and save_xlsx(f"{base}.xlsx", meta, jobs):
        saved_list.append(f"{base}.xlsx")
    if args.csv and save_csv(f"{base}.csv", jobs):
        saved_list.append(f"{base}.csv")

    if args.db:
        try:
            inserted = save_to_db(jobs, args.db)
            log.info("%d novas vagas salvas no banco %s", inserted, _get_db_path(args.db))
        except Exception as e:
            log.error("falha ao salvar no banco: %s", e)
            errors += 1

    log.info("%d vagas únicas encontradas.", len(jobs))
    if saved_list:
        log.info("Arquivos salvos: %s", ", ".join(saved_list))

    # Exit code: 2 se todas as queries falharam (útil em cronjobs)
    if plan and errors >= len(plan):
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
