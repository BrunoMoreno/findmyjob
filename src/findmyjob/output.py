"""Write search results to JSON, XLSX and CSV."""

from __future__ import annotations

import json
from pathlib import Path

from .console import paint

_CSV_FIELDS = ["title", "company", "link", "domain", "location", "query",
               "posted_at", "description"]


def save_json(path: str, meta: dict, jobs: list[dict]) -> None:
    payload = {**meta, "total": len(jobs), "jobs": jobs}
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                          encoding="utf-8")


def save_xlsx(path: str, meta: dict, jobs: list[dict]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(paint("[warning] openpyxl is not installed (pip install openpyxl); .xlsx was not generated.", "yellow"))
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    headers = ["#", "Title", "Company", "Link", "Domain", "Location", "Query", "Description"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center")

    for i, j in enumerate(jobs, start=1):
        ws.append([i, j.get("title", ""), j.get("company", ""), j.get("link", ""),
                   j.get("domain", ""), j.get("location", ""), j.get("query", ""),
                   j.get("description", "")])
        link_cell = ws.cell(row=i + 1, column=4)
        link_cell.hyperlink = j.get("link", "")
        link_cell.font = Font(color="0563C1", underline="single")

    # Dynamic widths based on content (with reasonable caps).
    caps = (5, 55, 28, 60, 22, 20, 45, 70)
    for idx, cap in enumerate(caps, start=1):
        longest = len(str(headers[idx - 1]))
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx,
                                values_only=True):
            if row[0]:
                longest = max(longest, len(str(row[0])))
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, cap)
    ws.freeze_panes = "A2"
    if jobs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(jobs) + 1}"

    info = wb.create_sheet("Search")
    for k, v in meta.items():
        info.append([k, ", ".join(v) if isinstance(v, list) else v])
    info.append(["total", len(jobs)])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 40

    wb.save(path)
    return True


def save_csv(path: str, jobs: list[dict]) -> bool:
    """Save jobs to CSV (UTF-8). Returns True on success."""
    import csv

    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=_CSV_FIELDS, extrasaction="ignore")
            writer.writeheader()
            for j in jobs:
                writer.writerow({k: j.get(k, "") or "" for k in _CSV_FIELDS})
    except OSError as e:
        print(paint(f"[warning] could not generate the CSV: {e}", "yellow"))
        return False
    return True
