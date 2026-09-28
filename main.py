#!/usr/bin/env python3
"""
Busca de vagas via dorks (site:dominio "cargo" "local").

Modo interativo (pergunta tudo):
    python job_search.py

Modo direto:
    python job_search.py "backend engineer" -l remote -l latam            # geral
    python job_search.py "backend engineer" -c br                         # só Brasil
    python job_search.py "backend engineer" -c br -l remote --group ats   # Brasil + remoto
    python job_search.py --list-countries

Dependências:
    pip install ddgs requests openpyxl

Backend google (opcional):
    export GOOGLE_API_KEY=...
    export GOOGLE_CX=...
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests

# Portais genéricos (funcionam para busca geral)
BOARDS = ["indeed.com", "linkedin.com/jobs", "glassdoor.com"]

# ATS usados pelas empresas (ouro escondido)
ATS = [
    "boards.greenhouse.io",
    "jobs.lever.co",
    "myworkdayjobs.com",
    "smartrecruiters.com",
    "jobs.ashbyhq.com",
    "apply.workable.com",
]

GROUPS = {"boards": BOARDS, "ats": ATS, "all": BOARDS + ATS}

# Config por país. Edite à vontade: "names" são os termos usados na query
# (com OR entre eles), "region" vai para o backend, "indeed"/"glassdoor" são
# os domínios locais e "extra" são portais que só existem naquele país.
COUNTRIES = {
    "br": {"names": ["Brazil", "Brasil"], "region": "br-pt", "gl": "br",
           "indeed": "br.indeed.com", "glassdoor": "glassdoor.com.br",
           "extra": ["gupy.io", "infojobs.com.br", "vagas.com.br"]},
    "pt": {"names": ["Portugal"], "region": "pt-pt", "gl": "pt",
           "indeed": "pt.indeed.com", "extra": []},
    "us": {"names": ["United States", "USA"], "region": "us-en", "gl": "us",
           "indeed": "indeed.com", "extra": []},
    "uk": {"names": ["United Kingdom", "UK"], "region": "uk-en", "gl": "uk",
           "indeed": "uk.indeed.com", "glassdoor": "glassdoor.co.uk", "extra": []},
    "ca": {"names": ["Canada"], "region": "ca-en", "gl": "ca",
           "indeed": "ca.indeed.com", "glassdoor": "glassdoor.ca", "extra": []},
    "de": {"names": ["Germany", "Deutschland"], "region": "de-de", "gl": "de",
           "indeed": "de.indeed.com", "glassdoor": "glassdoor.de", "extra": []},
    "es": {"names": ["Spain", "España"], "region": "es-es", "gl": "es",
           "indeed": "es.indeed.com", "glassdoor": "glassdoor.es", "extra": []},
    "mx": {"names": ["Mexico", "México"], "region": "mx-es", "gl": "mx",
           "indeed": "mx.indeed.com", "glassdoor": "glassdoor.com.mx", "extra": []},
    "ar": {"names": ["Argentina"], "region": "ar-es", "gl": "ar",
           "indeed": "ar.indeed.com", "glassdoor": "glassdoor.com.ar", "extra": []},
    "co": {"names": ["Colombia"], "region": "co-es", "gl": "co",
           "indeed": "co.indeed.com", "extra": []},
    "cl": {"names": ["Chile"], "region": "cl-es", "gl": "cl",
           "indeed": "cl.indeed.com", "extra": []},
}


# ----------------------------------------------------------------- queries --

def country_term(country: dict) -> str:
    names = [f'"{n}"' for n in country["names"]]
    return names[0] if len(names) == 1 else "(" + " OR ".join(names) + ")"


def domains_for(group: str, country: dict | None) -> list[str]:
    domains = list(GROUPS[group])
    if not country:
        return domains
    out = []
    for d in domains:
        if d == "indeed.com":
            d = country["indeed"]
        elif d == "glassdoor.com":
            d = country.get("glassdoor", "glassdoor.com")
        out.append(d)
    if group in ("boards", "all"):
        out += country["extra"]
    return out


def plan_queries(role: str, locals_: list[str], group: str,
                 country: dict | None) -> list[dict]:
    """Retorna [{domain, location, query}]."""
    plan = []
    for d in domains_for(group, country):
        if country:
            for r in (locals_ or [None]):
                terms = [country_term(country)] + ([f'"{r}"'] if r else [])
                label = country["names"][0] + (f" / {r}" if r else "")
                plan.append({"domain": d, "location": label,
                             "query": f'site:{d} "{role}" ' + " ".join(terms)})
        else:
            for loc in (locals_ or ["remote"]):
                plan.append({"domain": d, "location": loc,
                             "query": f'site:{d} "{role}" "{loc}"'})
    return plan


# ---------------------------------------------------------------- backends --

def search_ddg(query: str, max_results: int, country: dict | None) -> list[dict]:
    try:
        from ddgs import DDGS
    except ImportError:
        sys.exit("Instale com: pip install ddgs")
    kwargs = {"max_results": max_results}
    if country:
        kwargs["region"] = country["region"]
    results = DDGS().text(query, **kwargs)
    return [{"title": r.get("title", ""), "link": r.get("href", "")} for r in results]


def search_google(query: str, max_results: int, country: dict | None) -> list[dict]:
    key = os.environ.get("GOOGLE_API_KEY")
    cx = os.environ.get("GOOGLE_CX")
    if not key or not cx:
        sys.exit("Defina GOOGLE_API_KEY e GOOGLE_CX para usar o backend google.")
    params = {"key": key, "cx": cx, "q": query, "num": min(max_results, 10)}
    if country:
        params["gl"] = country["gl"]
    resp = requests.get("https://www.googleapis.com/customsearch/v1",
                        params=params, timeout=15)
    resp.raise_for_status()
    return [{"title": i.get("title", ""), "link": i.get("link", "")}
            for i in resp.json().get("items", [])]


BACKENDS = {"ddg": search_ddg, "google": search_google}


# ------------------------------------------------------------- interactive --

def ask(prompt: str, default: str = "") -> str:
    suffix = f" [{default}]" if default else ""
    answer = input(f"{prompt}{suffix}: ").strip()
    return answer or default


def default_basename(role: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", role.lower()).strip("-") or "vagas"
    return f"vagas_{slug}_{datetime.now():%Y%m%d_%H%M}"


def interactive_prompts(args: argparse.Namespace) -> None:
    print("=== Busca de vagas ===\n")
    while not args.role:
        args.role = ask("Qual vaga você está buscando? (ex: backend engineer)")

    print("\nAlcance da busca:")
    print("  1) Geral (mundo todo)")
    print("  2) Um país específico")
    if ask("Opção", "1") == "2":
        print("  Países: " + ", ".join(f"{k} ({v['names'][0]})" for k, v in COUNTRIES.items()))
        code = ask("Código do país", "br").lower()
        if code in COUNTRIES:
            args.country = code
        else:
            print(f"  [aviso] país '{code}' desconhecido, usando busca geral.")
        locs = ask("Refinar por cidade/remote (opcional, vírgula)", "")
    else:
        args.country = None
        locs = ask("Local(is), separados por vírgula", "remote, latam")
    args.local = [x.strip() for x in locs.split(",") if x.strip()]

    print("\nOnde buscar?")
    print("  1) Portais (Indeed, LinkedIn, Glassdoor + locais do país)")
    print("  2) ATS das empresas (Greenhouse, Lever, Workday, ...)")
    print("  3) Todos")
    args.group = {"1": "boards", "2": "ats", "3": "all"}.get(ask("Opção", "3"), "all")

    max_raw = ask("Resultados por query", str(args.max))
    args.max = int(max_raw) if max_raw.isdigit() else args.max

    args.output = ask("Nome base dos arquivos (gera .json e .xlsx)",
                      default_basename(args.role))
    print()


# ------------------------------------------------------------------ output --

def save_json(path: str, meta: dict, jobs: list[dict]) -> None:
    payload = {**meta, "total": len(jobs), "jobs": jobs}
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                          encoding="utf-8")


def save_xlsx(path: str, meta: dict, jobs: list[dict]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[aviso] openpyxl não instalado (pip install openpyxl); .xlsx não foi gerado.")
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Vagas"

    headers = ["#", "Título", "Link", "Domínio", "Local", "Query"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center")

    for i, j in enumerate(jobs, start=1):
        ws.append([i, j["title"], j["link"], j["domain"], j["location"], j["query"]])
        link_cell = ws.cell(row=i + 1, column=3)
        link_cell.hyperlink = j["link"]
        link_cell.font = Font(color="0563C1", underline="single")

    for col, width in zip("ABCDEF", (5, 60, 70, 24, 22, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    if jobs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(jobs) + 1}"

    info = wb.create_sheet("Busca")
    for k, v in meta.items():
        info.append([k, ", ".join(v) if isinstance(v, list) else v])
    info.append(["total", len(jobs)])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 40

    wb.save(path)
    return True


# -------------------------------------------------------------------- main --

def main() -> None:
    p = argparse.ArgumentParser(description="Busca de vagas com dorks de busca")
    p.add_argument("role", nargs="?", help='cargo, ex: "backend engineer" (omita para modo interativo)')
    p.add_argument("-c", "--country", choices=COUNTRIES,
                   help="limita a busca a um país (omita para busca geral)")
    p.add_argument("-l", "--local", action="append", default=[],
                   help="local/termo extra (repita para vários). Com -c refina o país (cidade, remote)")
    p.add_argument("-g", "--group", choices=GROUPS, default="all")
    p.add_argument("-b", "--backend", choices=BACKENDS, default="ddg")
    p.add_argument("-m", "--max", type=int, default=8, help="resultados por query")
    p.add_argument("-o", "--output", help="nome base dos arquivos de saída, gera .json e .xlsx")
    p.add_argument("-i", "--interactive", action="store_true", help="força o modo interativo")
    p.add_argument("--delay", type=float, default=2.0, help="segundos entre queries")
    p.add_argument("--show-queries", action="store_true", help="só imprime as queries")
    p.add_argument("--list-countries", action="store_true", help="lista os países disponíveis")
    args = p.parse_args()

    if args.list_countries:
        for k, v in COUNTRIES.items():
            print(f"{k}  {v['names'][0]}")
        return

    if args.interactive or not args.role:
        interactive_prompts(args)

    country = COUNTRIES.get(args.country) if args.country else None
    plan = plan_queries(args.role, args.local, args.group, country)

    if args.show_queries:
        for item in plan:
            print(item["query"])
        return

    scope = country["names"][0] if country else "geral"
    print(f"Escopo: {scope} | {len(plan)} queries")

    search = BACKENDS[args.backend]
    seen: set[str] = set()
    jobs: list[dict] = []

    try:
        for item in plan:
            print(f"\n== {item['domain']} | {item['location']} ==")
            print(f"   {item['query']}")
            try:
                results = search(item["query"], args.max, country)
            except Exception as e:  # rede, rate limit, chave inválida...
                print(f"   [erro] {e}")
                time.sleep(args.delay)
                continue

            new = [r for r in results if r["link"] and r["link"] not in seen]
            if not new:
                print("   (nenhum resultado novo)")
            for r in new:
                seen.add(r["link"])
                jobs.append({"title": r["title"], "link": r["link"],
                             "domain": item["domain"], "location": item["location"],
                             "query": item["query"]})
                print(f"   - {r['title']}\n     {r['link']}")
            time.sleep(args.delay)
    except KeyboardInterrupt:
        print("\n[interrompido] salvando o que foi encontrado até aqui...")

    base = args.output or default_basename(args.role)
    base = re.sub(r"\.(json|xlsx?)$", "", base, flags=re.IGNORECASE)
    meta = {
        "role": args.role,
        "scope": scope,
        "country": args.country or "",
        "locations": args.local,
        "group": args.group,
        "backend": args.backend,
        "searched_at": datetime.now().isoformat(timespec="seconds"),
    }
    save_json(f"{base}.json", meta, jobs)
    saved = [f"{base}.json"]
    if save_xlsx(f"{base}.xlsx", meta, jobs):
        saved.append(f"{base}.xlsx")
    print(f"\n{len(jobs)} vagas únicas encontradas. Salvo em: {', '.join(saved)}")


if __name__ == "__main__":
    main()